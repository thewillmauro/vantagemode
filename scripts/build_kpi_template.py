"""Builds downloads/gtm-kpi-scoreboard.xlsx, the KPI Scoreboard lead magnet.

KPI names, definitions, formulas, and services mirror the Scoreboard section
(#metrics) in index.html. Change them there first, then here, then rerun:
  python3 scripts/build_kpi_template.py   (needs openpyxl)
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path(__file__).resolve().parent.parent / "downloads" / "gtm-kpi-scoreboard.xlsx"

INK = "1A1712"
GOLD = "C9A84C"
GOLD_INK = "7A5B0B"
CREAM = "F5F0E8"
SAND = "ECE5D8"
MUTED = "5A5346"
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")
HEAD_FILL = PatternFill("solid", fgColor=INK)
BAND_FILL = PatternFill("solid", fgColor=SAND)
GOOD_FILL = PatternFill("solid", fgColor="D8EFD9")
BAD_FILL = PatternFill("solid", fgColor="F6D5D8")
THIN = Side(style="thin", color="D9D1C1")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

TITLE = Font(name="Arial", size=20, bold=True, color=INK)
H2 = Font(name="Arial", size=13, bold=True, color=GOLD_INK)
HEAD = Font(name="Arial", size=10, bold=True, color=CREAM)
BODY = Font(name="Arial", size=10, color=INK)
NOTE = Font(name="Arial", size=9, italic=True, color=MUTED)
BOLD = Font(name="Arial", size=10, bold=True, color=INK)
WRAP = Alignment(wrap_text=True, vertical="top")

SERVICES = {
    "01": "Data & Enrichment", "02": "AI Research", "03": "CRM & RevOps", "04": "Outbound & Lifecycle",
    "05": "Product-Led Growth", "06": "Conversation Intelligence",
}

# (group, lever, name, definition, formula, services, cadence, better)
KPIS = [
    ("Pipeline", "More qualified at-bats", "Qualified pipeline created", "Dollar value of new opportunities that match your ideal customer profile (ICP), per month.", "Sum of new opportunity value", "01 02 04 05", "Monthly", "Higher"),
    ("Pipeline", "More qualified at-bats", "Product-qualified accounts", "Accounts whose product usage crosses your buying threshold.", "Accounts over the usage threshold ÷ active accounts", "05", "Monthly", "Higher"),
    ("Pipeline", "More qualified at-bats", "Enrichment match rate", "Share of accounts with every required field filled.", "Complete records ÷ total records", "01", "Weekly", "Higher"),
    ("Pipeline", "More qualified at-bats", "Decision-maker accuracy", "Share of researched accounts where the contact is the real buyer, graded against hand-labeled accounts.", "Correct contacts ÷ labeled sample", "02", "Monthly", "Higher"),
    ("Conversion", "Win more of them", "Speed to lead", "Time from a signup or form fill to the first human touch.", "Median minutes, inbound only", "03 04", "Weekly", "Lower"),
    ("Conversion", "Win more of them", "Positive reply rate", "Replies that want to keep talking, per 100 delivered emails.", "Positive replies ÷ delivered", "04", "Weekly", "Higher"),
    ("Conversion", "Win more of them", "Meeting rate", "Meetings held per 100 accounts contacted.", "Meetings held ÷ accounts contacted", "03 04", "Weekly", "Higher"),
    ("Conversion", "Win more of them", "Win rate", "Share of opportunities that close, and how they move stage to stage.", "Won ÷ opportunities created", "03 06", "Quarterly", "Higher"),
    ("Conversion", "Win more of them", "Free-to-paid conversion", "Share of free accounts that upgrade in a period.", "Upgrades ÷ free accounts at start", "05", "Monthly", "Higher"),
    ("Deal Size", "Earn more per customer", "Average contract value", "Average first-year value of a new customer. ARR is annual recurring revenue.", "New ARR ÷ new customers", "03 05", "Quarterly", "Higher"),
    ("Deal Size", "Earn more per customer", "Expansion pipeline", "Pipeline from more seats, higher tiers, or new teams at existing customers.", "Sum of expansion opportunity value", "05 06", "Quarterly", "Higher"),
    ("Retention", "Keep and grow them", "Net revenue retention", "Revenue kept and grown from last year's customers, after churn and expansion.", "(Start ARR + expansion − churn − contraction) ÷ start ARR", "05 06", "Quarterly", "Higher"),
    ("Retention", "Keep and grow them", "Activation rate", "Share of new signups that reach their first moment of real value.", "Activated signups ÷ total signups", "05", "Monthly", "Higher"),
    ("Retention", "Keep and grow them", "Churn signals actioned", "Share of at-risk flags from usage or calls that got a response within a week.", "Flags actioned in 7 days ÷ flags raised", "03 06", "Monthly", "Higher"),
    ("Efficiency", "Growth that pays for itself", "Cost per qualified account", "Data, AI, and tool spend for every account that reaches Scored.", "Pipeline spend ÷ accounts scored", "01 02", "Monthly", "Lower"),
    ("Efficiency", "Growth that pays for itself", "CAC payback", "Customer acquisition cost (CAC) is what it costs to win a customer. Payback is the months of gross margin it takes to earn that back.", "CAC ÷ (monthly revenue per customer × gross margin)", "01 03 04", "Quarterly", "Lower"),
    ("Efficiency", "Growth that pays for itself", "LTV to CAC", "Lifetime value (LTV) is the gross margin a customer brings over their whole time with you, compared with what it cost to win them.", "(Monthly revenue per customer × gross margin ÷ monthly churn rate) ÷ CAC", "03 05 06", "Quarterly", "Higher"),
    ("Efficiency", "Growth that pays for itself", "Rep selling time", "Share of rep hours spent with buyers instead of research and list building.", "Hours with buyers ÷ total rep hours", "02 04", "Weekly", "Higher"),
    ("System Health", "Numbers you can trust", "Data freshness", "Days since each account was last enriched.", "Median days since last enrichment", "01", "Weekly", "Lower"),
    ("System Health", "Numbers you can trust", "Duplicate rate", "Duplicate records a sync creates. The target is zero.", "Duplicates created ÷ records synced", "03", "Weekly", "Lower"),
    ("System Health", "Numbers you can trust", "Deliverability", "Share of email that lands in the inbox, plus bounce and spam-complaint rates.", "Delivered to inbox ÷ sent", "04", "Weekly", "Higher"),
    ("System Health", "Numbers you can trust", "Draft approval rate", "Share of AI drafts approved without edits. When it falls, the inputs need work.", "Approved as written ÷ drafts reviewed", "02 04", "Weekly", "Higher"),
]


def services(codes):
    return ", ".join(f"{c} {SERVICES[c]}" for c in codes.split())


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, Alignment(vertical="center", wrap_text=True), BOX
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = sub
    ws["A2"].font = NOTE
    ws.row_dimensions[1].height = 30


def period_grid(ws, start_row, periods, period_label, inputs, kpis):
    """Inputs block (yellow cells) then a KPI block computed from it, with target + status columns.

    inputs: [(key, label, number_format)]
    kpis:   [(name, excel_formula_template using {key} -> cell ref, number_format, better)]
    """
    first = 4  # first period column (D); A = name, B = target, C = better
    cols = [get_column_letter(first + i) for i in range(periods)]
    r = start_row
    header(ws, r, ["Inputs", "", "Unit"] + [f"{period_label} {i + 1}" for i in range(periods)])
    refs = {}
    for key, label, fmt in inputs:
        r += 1
        ws.cell(row=r, column=1, value=label).font = BODY
        ws.cell(row=r, column=3, value={"0": "#", "$#,##0": "$", "0.0": "min/days"}.get(fmt, "")).font = NOTE
        for col in cols:
            c = ws[f"{col}{r}"]
            c.fill, c.border, c.number_format = INPUT_FILL, BOX, fmt
        refs[key] = r
    r += 2
    header(ws, r, ["KPI", "Your target", "Better"] + [f"{period_label} {i + 1}" for i in range(periods)])
    for name, tmpl, fmt, better in kpis:
        r += 1
        ws.cell(row=r, column=1, value=name).font = BOLD
        t = ws.cell(row=r, column=2)
        t.fill, t.border, t.number_format = INPUT_FILL, BOX, fmt
        ws.cell(row=r, column=3, value=better).font = NOTE
        for col in cols:
            expr = tmpl.format(**{k: f"{col}{v}" for k, v in refs.items()})
            c = ws[f"{col}{r}"]
            c.value = f'=IFERROR({expr},"")'
            c.number_format, c.border, c.font = fmt, BOX, BODY
        rng = f"{cols[0]}{r}:{cols[-1]}{r}"
        tl = f"{cols[0]}{r}"
        good = f'AND({tl}<>"",$B{r}<>"",IF($C{r}="Higher",{tl}>=$B{r},{tl}<=$B{r}))'
        bad = f'AND({tl}<>"",$B{r}<>"",IF($C{r}="Higher",{tl}<$B{r},{tl}>$B{r}))'
        ws.conditional_formatting.add(rng, FormulaRule(formula=[good], fill=GOOD_FILL))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[bad], fill=BAD_FILL))
    return r


def build():
    wb = Workbook()

    # ── Start Here ──
    ws = wb.active
    ws.title = "Start Here"
    title(ws, "GTM KPI Scoreboard", "From Vantage Mode · vantagemode.com · Free to use and adapt")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 90
    rows = [
        ("How to use it", None),
        ("1. Pick five", "Do not start with all 22 metrics. Pick the five that match your biggest gap. If you are not sure, start with Qualified pipeline created, Speed to lead, Meeting rate, Win rate, and Rep selling time."),
        ("2. Baseline", "Fill in the yellow cells on the Weekly tab for the last four weeks you can reconstruct. That is your baseline. Growth is measured against your numbers, not someone else's benchmark."),
        ("3. Set targets", "Enter a target for each KPI in the Your target column. Green means on target, red means off. There are no default targets on purpose."),
        ("4. Review on a cadence", "Weekly for activity and data health, monthly for pipeline and efficiency, quarterly for win rate, contract value, and retention. See the table below."),
        ("5. Test, do not guess", "Every change is an experiment. Write the hypothesis and the decision rule in the Experiment Log before you start, so the result decides."),
        ("", None),
        ("The tabs", None),
        ("Weekly", "13 weeks of inputs and the weekly KPIs computed from them."),
        ("Monthly", "12 months of inputs and the monthly KPIs."),
        ("Quarterly", "4 quarters of inputs and the quarterly KPIs, including CAC payback and LTV to CAC."),
        ("KPI Definitions", "All 22 metrics: what each one means, how it is calculated, and what moves it."),
        ("Experiment Log", "Hypothesis, test, metric, decision rule, result, and decision for every change."),
        ("", None),
        ("Review cadence", None),
        ("Weekly", "Review reply and meeting rates, speed to lead, deliverability, and data health. Decide how to tune lists, copy, and routing."),
        ("Monthly", "Review pipeline created, stage conversion, product-qualified accounts, and cost per qualified account. Decide how to reweight scoring, shift channels, and fund what is working."),
        ("Quarterly", "Review win rate by segment, contract value, net revenue retention, and CAC payback. Revisit the ICP, the roadmap, and where the next build goes."),
        ("", None),
        ("Legend", None),
        ("Yellow cells", "Your inputs. Everything else is calculated."),
        ("Green / red", "On or off the target you set."),
        ("", None),
        ("Want help?", "We build the systems that move these numbers. Take the free GTM Readiness Scorecard at vantagemode.com/scorecard, or book a free 20-minute Pipeline Teardown at vantagemode.com."),
    ]
    r = 4
    for a, b in rows:
        ca = ws.cell(row=r, column=1, value=a)
        if b is None:
            ca.font = H2
        else:
            ca.font = BOLD
            cb = ws.cell(row=r, column=2, value=b)
            cb.font, cb.alignment = BODY, WRAP
            ca.alignment = WRAP
            if a == "Yellow cells":
                ca.fill = INPUT_FILL
        r += 1
    ws.sheet_view.showGridLines = False

    # ── Weekly ──
    ws = wb.create_sheet("Weekly")
    title(ws, "Weekly Scoreboard", "Yellow cells are inputs. Enter the Monday of week 1 in D4; the other dates fill in.")
    ws["A4"] = "Week starting"
    ws["A4"].font = BOLD
    ws["D4"].fill, ws["D4"].border, ws["D4"].number_format = INPUT_FILL, BOX, "mmm d"
    for i in range(1, 13):
        col, prev = get_column_letter(4 + i), get_column_letter(3 + i)
        ws[f"{col}4"] = f'=IF({prev}4="","",{prev}4+7)'
        ws[f"{col}4"].number_format = "mmm d"
    weekly_inputs = [
        ("opp_value", "New qualified opportunity value", "$#,##0"),
        ("accounts", "Accounts contacted", "0"),
        ("meetings", "Meetings held", "0"),
        ("sent", "Emails sent", "0"),
        ("delivered", "Emails delivered", "0"),
        ("inbox", "Emails landing in the inbox", "0"),
        ("positive", "Positive replies", "0"),
        ("stl", "Median minutes to first touch (inbound)", "0.0"),
        ("records", "Account records in the CRM", "0"),
        ("complete", "Records with every required field", "0"),
        ("fresh", "Median days since last enrichment", "0.0"),
        ("synced", "Records synced", "0"),
        ("dupes", "Duplicates created", "0"),
        ("drafts", "AI drafts reviewed", "0"),
        ("approved", "Drafts approved as written", "0"),
        ("rep_hours", "Total rep hours", "0"),
        ("buyer_hours", "Rep hours with buyers", "0"),
    ]
    weekly_kpis = [
        ("Qualified pipeline created", "IF({opp_value}=\"\",NA(),{opp_value})", "$#,##0", "Higher"),
        ("Speed to lead (minutes)", "IF({stl}=\"\",NA(),{stl})", "0.0", "Lower"),
        ("Positive reply rate", "{positive}/{delivered}", "0.0%", "Higher"),
        ("Meeting rate", "{meetings}/{accounts}", "0.0%", "Higher"),
        ("Deliverability", "{inbox}/{sent}", "0.0%", "Higher"),
        ("Enrichment match rate", "{complete}/{records}", "0.0%", "Higher"),
        ("Data freshness (days)", "IF({fresh}=\"\",NA(),{fresh})", "0.0", "Lower"),
        ("Duplicate rate", "{dupes}/{synced}", "0.0%", "Lower"),
        ("Draft approval rate", "{approved}/{drafts}", "0.0%", "Higher"),
        ("Rep selling time", "{buyer_hours}/{rep_hours}", "0.0%", "Higher"),
    ]
    period_grid(ws, 6, 13, "Wk", weekly_inputs, weekly_kpis)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 9
    for i in range(13):
        ws.column_dimensions[get_column_letter(4 + i)].width = 11
    ws.freeze_panes = "D5"

    # ── Monthly ──
    ws = wb.create_sheet("Monthly")
    title(ws, "Monthly Scoreboard", "Yellow cells are inputs. Pipeline created can be the sum of the four or five weeks in that month.")
    monthly_inputs = [
        ("opp_value", "New qualified opportunity value", "$#,##0"),
        ("active", "Active product accounts", "0"),
        ("pqa", "Accounts over the usage threshold", "0"),
        ("labeled", "Researched accounts in the labeled sample", "0"),
        ("correct", "Sample accounts with the right buyer", "0"),
        ("free_start", "Free accounts at start of month", "0"),
        ("upgrades", "Upgrades to paid", "0"),
        ("signups", "New signups", "0"),
        ("activated", "Signups that reached first value", "0"),
        ("flags", "At-risk flags raised", "0"),
        ("flags_done", "Flags actioned within 7 days", "0"),
        ("spend", "Pipeline spend (data, AI, tools)", "$#,##0"),
        ("scored", "Accounts that reached Scored", "0"),
    ]
    monthly_kpis = [
        ("Qualified pipeline created", "IF({opp_value}=\"\",NA(),{opp_value})", "$#,##0", "Higher"),
        ("Product-qualified accounts", "{pqa}/{active}", "0.0%", "Higher"),
        ("Decision-maker accuracy", "{correct}/{labeled}", "0.0%", "Higher"),
        ("Free-to-paid conversion", "{upgrades}/{free_start}", "0.0%", "Higher"),
        ("Activation rate", "{activated}/{signups}", "0.0%", "Higher"),
        ("Churn signals actioned", "{flags_done}/{flags}", "0.0%", "Higher"),
        ("Cost per qualified account", "{spend}/{scored}", "$#,##0.00", "Lower"),
    ]
    period_grid(ws, 4, 12, "Mo", monthly_inputs, monthly_kpis)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 9
    for i in range(12):
        ws.column_dimensions[get_column_letter(4 + i)].width = 11
    ws.freeze_panes = "D5"

    # ── Quarterly ──
    ws = wb.create_sheet("Quarterly")
    title(ws, "Quarterly Scoreboard", "Yellow cells are inputs. Use a 12-month window for net revenue retention.")
    quarterly_inputs = [
        ("opps", "Opportunities created", "0"),
        ("won", "Opportunities won", "0"),
        ("new_arr", "New ARR won", "$#,##0"),
        ("new_cust", "New customers", "0"),
        ("expansion_pipe", "Expansion opportunity value", "$#,##0"),
        ("start_arr", "ARR from these customers 12 months ago", "$#,##0"),
        ("expansion", "Expansion ARR", "$#,##0"),
        ("churn", "Churned ARR", "$#,##0"),
        ("contraction", "Contraction ARR", "$#,##0"),
        ("sm_spend", "Sales and marketing spend", "$#,##0"),
        ("mrr_cust", "Monthly revenue per customer", "$#,##0"),
        ("margin", "Gross margin (%)", "0%"),
        ("churn_rate", "Monthly customer churn rate (%)", "0.0%"),
    ]
    quarterly_kpis = [
        ("Win rate", "{won}/{opps}", "0.0%", "Higher"),
        ("Average contract value", "{new_arr}/{new_cust}", "$#,##0", "Higher"),
        ("Expansion pipeline", "IF({expansion_pipe}=\"\",NA(),{expansion_pipe})", "$#,##0", "Higher"),
        ("Net revenue retention", "({start_arr}+{expansion}-{churn}-{contraction})/{start_arr}", "0.0%", "Higher"),
        ("CAC", "{sm_spend}/{new_cust}", "$#,##0", "Lower"),
        ("CAC payback (months)", "({sm_spend}/{new_cust})/({mrr_cust}*{margin})", "0.0", "Lower"),
        ("LTV to CAC", "({mrr_cust}*{margin}/{churn_rate})/({sm_spend}/{new_cust})", "0.0", "Higher"),
    ]
    period_grid(ws, 4, 4, "Q", quarterly_inputs, quarterly_kpis)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 13
    ws.column_dimensions["C"].width = 9
    for i in range(4):
        ws.column_dimensions[get_column_letter(4 + i)].width = 14
    ws.freeze_panes = "D5"

    # ── KPI Definitions ──
    ws = wb.create_sheet("KPI Definitions")
    title(ws, "KPI Definitions", "Revenue growth = Pipeline × Conversion × Deal Size × Retention, kept efficient and trustworthy by Efficiency and System Health.")
    header(ws, 4, ["Group", "Lever", "KPI", "Definition", "Formula", "Review", "Better", "What moves it"], [14, 24, 28, 60, 44, 11, 9, 40])
    for i, (group, lever, name, definition, formula, svc, cadence, better) in enumerate(KPIS):
        r = 5 + i
        for col, val in enumerate([group, lever, name, definition, formula, cadence, better, services(svc)], 1):
            c = ws.cell(row=r, column=col, value=val)
            c.font = BOLD if col == 3 else BODY
            c.alignment, c.border = WRAP, BOX
            if i % 2:
                c.fill = BAND_FILL
    ws.freeze_panes = "D5"
    ws.auto_filter.ref = f"A4:H{4 + len(KPIS)}"

    # ── Experiment Log ──
    ws = wb.create_sheet("Experiment Log")
    title(ws, "Experiment Log", "Agree the decision rule before the test starts, so the result decides, not opinions. Log every result, including the losers.")
    cols = ["#", "Hypothesis", "Test", "Metric", "Decision rule", "Start", "End", "Result", "Decision", "Notes"]
    header(ws, 4, cols, [5, 40, 36, 30, 36, 11, 11, 30, 12, 30])
    example = [1, "Active free-plan users answer text messages faster than email.", "Split 100 active users evenly between the two channels for two weeks.",
               "Meetings per 100 contacts, opt-outs, replies that want to keep talking", "Keep the channel that books more meetings without raising opt-outs.", None, None, "", "", "Example row. Replace with your own."]
    dv = DataValidation(type="list", formula1='"Keep,Kill,Extend,Rerun"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(5, 41):
        for col in range(1, len(cols) + 1):
            c = ws.cell(row=r, column=col)
            c.border, c.alignment, c.font = BOX, WRAP, BODY
            if col in (6, 7):
                c.number_format = "yyyy-mm-dd"
        dv.add(f"I{r}")
    for col, val in enumerate(example, 1):
        ws.cell(row=5, column=col, value=val)
    ws.cell(row=5, column=10).font = NOTE
    ws.freeze_panes = "C5"

    for sheet in wb.worksheets:
        sheet.sheet_properties.tabColor = GOLD
    wb.properties.title = "GTM KPI Scoreboard"
    wb.properties.creator = "Vantage Mode"
    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    print(f"wrote {OUT} ({len(KPIS)} KPIs)")


if __name__ == "__main__":
    build()
